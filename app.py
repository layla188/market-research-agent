"""
Streamlit interface for the Autonomous Market Research Agent.

Run with:
    streamlit run app.py
"""

from __future__ import annotations

import io
import json
import re
import uuid
from typing import Any

import streamlit as st
from langgraph.types import Command

from src.graph import research_graph


# -------------------------------------------------------------------
# Page configuration
# -------------------------------------------------------------------

st.set_page_config(
    page_title="Autonomous Market Research Agent",
    page_icon="🔎",
    layout="wide",
    initial_sidebar_state="expanded",
)


# -------------------------------------------------------------------
# Session state
# -------------------------------------------------------------------

DEFAULT_SESSION_VALUES = {
    "thread_id": None,
    "topic": "",
    "result": None,
    "phase": "idle",  # idle | approval | complete | error
    "error": "",
    "pending_report": "",
    "final_report": "",
    "interrupt_payload": None,
}


def initialize_session() -> None:
    for key, value in DEFAULT_SESSION_VALUES.items():
        if key not in st.session_state:
            st.session_state[key] = value

    if st.session_state.thread_id is None:
        st.session_state.thread_id = str(uuid.uuid4())


def reset_session() -> None:
    st.session_state.thread_id = str(uuid.uuid4())
    st.session_state.topic = ""
    st.session_state.result = None
    st.session_state.phase = "idle"
    st.session_state.error = ""
    st.session_state.pending_report = ""
    st.session_state.final_report = ""
    st.session_state.interrupt_payload = None


initialize_session()


# -------------------------------------------------------------------
# Helpers
# -------------------------------------------------------------------

def build_initial_state(topic: str) -> dict[str, Any]:
    return {
        "topic": topic,
        "research_plan": [],
        "research_results": [],
        "evidence_status": "",
        "research_gaps": [],
        "research_iteration": 0,
        "analysis": "",
        "draft_report": "",
        "critique": "",
        "review_status": "",
        "iteration": 0,
        "approved": False,
        "memory_notes": [],
        "errors": [],
    }


def graph_config() -> dict[str, Any]:
    return {
        "configurable": {
            "thread_id": st.session_state.thread_id,
        }
    }


def normalize_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    if isinstance(value, list):
        parts: list[str] = []

        for item in value:
            if isinstance(item, str):
                parts.append(item)

            elif isinstance(item, dict):
                text_value = (
                    item.get("text")
                    or item.get("content")
                    or item.get("value")
                )

                if isinstance(text_value, str):
                    parts.append(text_value)

        if parts:
            return "\n\n".join(
                part.strip()
                for part in parts
                if part.strip()
            )

        return json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
        )

    if isinstance(value, dict):
        for key in ("text", "content", "value"):
            text_value = value.get(key)

            if isinstance(text_value, str):
                return text_value.strip()

        return json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
        )

    return str(value).strip()


def looks_like_report(text: str) -> bool:
    clean = normalize_text(text)

    if len(clean) < 300:
        return False

    report_markers = (
        "Executive Summary",
        "Market Overview",
        "Competitive",
        "Sources",
    )

    return any(
        marker.lower() in clean.lower()
        for marker in report_markers
    )


def get_interrupt_payload(
    result: dict[str, Any] | None,
) -> dict[str, Any] | None:

    if not result:
        return None

    interrupts = result.get("__interrupt__")

    if not interrupts:
        return None

    first_interrupt = interrupts[0]

    if hasattr(first_interrupt, "value"):
        payload = first_interrupt.value

    elif isinstance(first_interrupt, dict):
        payload = first_interrupt

    else:
        return None

    if isinstance(payload, dict):
        return payload

    return {"message": str(payload)}


def get_checkpoint_state() -> dict[str, Any]:
    try:
        snapshot = research_graph.get_state(
            graph_config()
        )

        if snapshot is None:
            return {}

        values = getattr(snapshot, "values", None)

        if isinstance(values, dict):
            return dict(values)

    except Exception:
        pass

    return {}


def merge_states(*states: Any) -> dict[str, Any]:
    merged: dict[str, Any] = {}

    for state in states:
        if isinstance(state, dict):
            merged.update(state)

    return merged


def choose_report(*candidates: Any) -> str:
    normalized = [
        normalize_text(value)
        for value in candidates
    ]

    for text in normalized:
        if looks_like_report(text):
            return text

    non_empty = [
        text
        for text in normalized
        if text
    ]

    if non_empty:
        return max(
            non_empty,
            key=len,
        )

    return ""


def safe_filename(topic: str) -> str:
    filename = topic.lower().strip()

    filename = re.sub(
        r"[^a-z0-9]+",
        "-",
        filename,
    )

    filename = filename.strip("-")

    return (
        filename[:60]
        or "market-research-report"
    )


# -------------------------------------------------------------------
# Graph execution
# -------------------------------------------------------------------

def run_research(topic: str) -> None:
    try:
        st.session_state.error = ""
        st.session_state.topic = topic
        st.session_state.pending_report = ""
        st.session_state.final_report = ""
        st.session_state.interrupt_payload = None

        invoke_result = research_graph.invoke(
            build_initial_state(topic),
            config=graph_config(),
        )

        checkpoint_state = get_checkpoint_state()

        result = merge_states(
            invoke_result,
            checkpoint_state,
        )

        payload = get_interrupt_payload(
            invoke_result
        )

        if payload:

            reviewed_report = choose_report(
                payload.get("report"),
                result.get("draft_report"),
                invoke_result.get("draft_report")
                if isinstance(invoke_result, dict)
                else "",
            )

            st.session_state.pending_report = (
                reviewed_report
            )

            st.session_state.interrupt_payload = (
                payload
            )

            st.session_state.result = result
            st.session_state.phase = "approval"

        else:

            final_report = choose_report(
                result.get("draft_report"),
                invoke_result.get("draft_report")
                if isinstance(invoke_result, dict)
                else "",
            )

            st.session_state.final_report = (
                final_report
            )

            st.session_state.result = result
            st.session_state.phase = "complete"

    except Exception as exc:
        st.session_state.error = str(exc)
        st.session_state.phase = "error"


def resume_after_human_decision(
    decision: str,
) -> None:

    try:
        st.session_state.error = ""

        before_resume = (
            st.session_state.result or {}
        )

        reviewed_report = (
            st.session_state.pending_report
        )

        resume_result = research_graph.invoke(
            Command(resume=decision),
            config=graph_config(),
        )

        checkpoint_state = get_checkpoint_state()

        final_state = merge_states(
            before_resume,
            resume_result,
            checkpoint_state,
        )

        final_report = choose_report(
            reviewed_report,
            final_state.get("draft_report"),
            before_resume.get("draft_report")
            if isinstance(before_resume, dict)
            else "",
        )

        if final_report:
            final_state["draft_report"] = (
                final_report
            )

        st.session_state.final_report = (
            final_report
        )

        st.session_state.result = final_state
        st.session_state.interrupt_payload = None
        st.session_state.phase = "complete"

    except Exception as exc:
        st.session_state.error = str(exc)
        st.session_state.phase = "error"


# -------------------------------------------------------------------
# Report export
# -------------------------------------------------------------------

def markdown_to_plain_text(
    markdown_text: str,
) -> str:

    text = markdown_text

    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text,
    )

    text = re.sub(
        r"\*(.*?)\*",
        r"\1",
        text,
    )

    text = re.sub(
        r"`(.*?)`",
        r"\1",
        text,
    )

    return text


def create_docx(report_text: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()

    # Title
    title = document.add_heading(
        "Market Research Report",
        level=0,
    )

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(
        st.session_state.topic
    )

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    document.add_paragraph("")

    for line in report_text.splitlines():

        stripped = line.strip()

        if not stripped:
            document.add_paragraph("")
            continue

        # Markdown headings
        if stripped.startswith("# "):
            document.add_heading(
                markdown_to_plain_text(
                    stripped[2:]
                ),
                level=1,
            )

        elif stripped.startswith("## "):
            document.add_heading(
                markdown_to_plain_text(
                    stripped[3:]
                ),
                level=2,
            )

        elif stripped.startswith("### "):
            document.add_heading(
                markdown_to_plain_text(
                    stripped[4:]
                ),
                level=3,
            )

        # Bullet points
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                markdown_to_plain_text(
                    stripped[2:]
                )
            )

        # Numbered lists
        elif re.match(
            r"^\d+\.\s+",
            stripped,
        ):
            clean = re.sub(
                r"^\d+\.\s+",
                "",
                stripped,
            )

            paragraph = document.add_paragraph(
                style="List Number"
            )

            paragraph.add_run(
                markdown_to_plain_text(
                    clean
                )
            )

        else:
            document.add_paragraph(
                markdown_to_plain_text(
                    stripped
                )
            )

    # Basic font
    styles = document.styles

    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


def create_pdf(report_text: str) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import (
        ParagraphStyle,
        getSampleStyleSheet,
    )
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Market Research Report",
        author="Autonomous Market Research Agent",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=20,
        leading=24,
        spaceAfter=8,
    )

    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=11,
        leading=14,
        spaceAfter=18,
    )

    heading1_style = ParagraphStyle(
        "ReportHeading1",
        parent=styles["Heading1"],
        fontSize=15,
        leading=19,
        spaceBefore=12,
        spaceAfter=8,
    )

    heading2_style = ParagraphStyle(
        "ReportHeading2",
        parent=styles["Heading2"],
        fontSize=12.5,
        leading=16,
        spaceBefore=10,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontSize=9.5,
        leading=14,
        spaceAfter=7,
    )

    bullet_style = ParagraphStyle(
        "ReportBullet",
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-7,
        spaceAfter=4,
    )

    story = []

    story.append(
        Paragraph(
            "Market Research Report",
            title_style,
        )
    )

    story.append(
        Paragraph(
            markdown_to_plain_text(
                st.session_state.topic
            ),
            subtitle_style,
        )
    )

    for line in report_text.splitlines():

        stripped = line.strip()

        if not stripped:
            story.append(
                Spacer(1, 4)
            )
            continue

        # Escape XML-sensitive characters
        clean = (
            markdown_to_plain_text(
                stripped
            )
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if stripped.startswith("# "):
            story.append(
                Paragraph(
                    clean[2:],
                    heading1_style,
                )
            )

        elif stripped.startswith("## "):
            story.append(
                Paragraph(
                    clean[3:],
                    heading1_style,
                )
            )

        elif stripped.startswith("### "):
            story.append(
                Paragraph(
                    clean[4:],
                    heading2_style,
                )
            )

        elif stripped.startswith("- "):
            story.append(
                Paragraph(
                    "• " + clean[2:],
                    bullet_style,
                )
            )

        elif re.match(
            r"^\d+\.\s+",
            stripped,
        ):
            story.append(
                Paragraph(
                    clean,
                    bullet_style,
                )
            )

        else:
            story.append(
                Paragraph(
                    clean,
                    body_style,
                )
            )

    document.build(story)

    return output.getvalue()


# -------------------------------------------------------------------
# Workflow visualization
# -------------------------------------------------------------------

def render_workflow(
    phase: str,
    result: dict[str, Any],
) -> None:

    st.subheader("Research workflow")

    if phase == "idle":
        current_index = 0

    elif phase == "approval":
        current_index = 5

    elif phase == "complete":
        current_index = 6

    else:
        current_index = 0

    stages = [
        "Planning",
        "Research",
        "Evidence analysis",
        "Report writing",
        "Review",
        "Human approval",
    ]

    cols = st.columns(len(stages))

    for index, (col, stage) in enumerate(
        zip(cols, stages)
    ):

        if phase == "complete":
            symbol = "✓"
            state = "complete"

        elif phase == "approval":
            if index < 5:
                symbol = "✓"
                state = "complete"
            elif index == 5:
                symbol = "●"
                state = "current"
            else:
                symbol = "○"
                state = "pending"

        elif phase == "error":
            symbol = "!"
            state = "error"

        elif index < current_index:
            symbol = "✓"
            state = "complete"

        elif index == current_index:
            symbol = "●"
            state = "current"

        else:
            symbol = "○"
            state = "pending"

        with col:

            if state == "complete":
                st.success(
                    f"{symbol} {stage}"
                )

            elif state == "current":
                st.info(
                    f"{symbol} {stage}"
                )

            elif state == "error":
                st.error(
                    f"{symbol} {stage}"
                )

            else:
                st.caption(
                    f"{symbol} {stage}"
                )


# -------------------------------------------------------------------
# Status / dashboard
# -------------------------------------------------------------------

def render_status(
    result: dict[str, Any],
) -> None:

    st.subheader("Research overview")

    evidence = (
        result.get("evidence_status")
        or "Not evaluated"
    )

    research_iterations = result.get(
        "research_iteration",
        0,
    )

    revisions = result.get(
        "iteration",
        0,
    )

    reviewer_status = (
        result.get("review_status")
        or "Not reviewed"
    )

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric(
            "Evidence",
            evidence,
        )

    with col2:
        st.metric(
            "Research iterations",
            research_iterations,
        )

    with col3:
        st.metric(
            "Report revisions",
            revisions,
        )

    with col4:
        st.metric(
            "Reviewer status",
            reviewer_status,
        )

    gaps = result.get(
        "research_gaps",
        [],
    )

    errors = result.get(
        "errors",
        [],
    )

    if gaps:
        with st.expander(
            f"Research gaps ({len(gaps)})"
        ):
            for gap in gaps:
                st.write(
                    f"• {gap}"
                )

    if errors:
        with st.expander(
            f"Workflow warnings ({len(errors)})",
            expanded=True,
        ):
            for error in errors:
                st.warning(error)


def render_report(
    report: Any,
    heading: str = "Research report",
) -> None:

    report_text = normalize_text(report)

    st.subheader(heading)

    if not report_text:
        st.warning(
            "No report text is available."
        )
        return

    if not looks_like_report(report_text):
        st.warning(
            "The generated report appears incomplete."
        )

    with st.container(border=True):
        st.markdown(report_text)

    filename_topic = safe_filename(
        st.session_state.topic
    )

    st.markdown("### Export report")

    pdf_col, docx_col = st.columns(2)

    with pdf_col:
        try:
            pdf_data = create_pdf(
                report_text
            )

            st.download_button(
                label="Download PDF",
                data=pdf_data,
                file_name=(
                    f"{filename_topic}.pdf"
                ),
                mime="application/pdf",
                type="primary",
                use_container_width=True,
                icon=":material/picture_as_pdf:",
                on_click="ignore",
            )

        except ImportError:
            st.error(
                "PDF export requires reportlab. "
                "Install it with: pip install reportlab"
            )

    with docx_col:
        try:
            docx_data = create_docx(
                report_text
            )

            st.download_button(
                label="Download Word",
                data=docx_data,
                file_name=(
                    f"{filename_topic}.docx"
                ),
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                ),
                use_container_width=True,
                icon=":material/description:",
                 on_click="ignore",
            )

        except ImportError:
            st.error(
                "Word export requires python-docx. "
                "Install it with: pip install python-docx"
            )


def render_memory(
    result: dict[str, Any],
) -> None:

    notes = result.get(
        "memory_notes",
        [],
    )

    with st.expander(
        "Research memory",
        expanded=False,
    ):

        if not notes:
            st.caption(
                "No research-memory entries "
                "are present in the final state."
            )
            return

        for index, note in enumerate(
            notes,
            start=1,
        ):

            claim = note.get(
                "claim",
                "",
            )

            source = note.get(
                "source",
                "",
            )

            topic = note.get(
                "topic",
                "",
            )

            date = note.get(
                "date",
                "",
            )

            st.markdown(
                f"**Finding {index}**"
            )

            if claim:
                st.write(claim)

            if source:
                st.caption(
                    f"Source: {source}"
                )

            if date:
                st.caption(
                    f"Date: {date}"
                )

            if topic:
                st.caption(
                    f"Topic: {topic}"
                )

            if index < len(notes):
                st.divider()


def render_developer_details(
    result: dict[str, Any],
) -> None:

    with st.expander(
        "Developer details",
        expanded=False,
    ):

        st.caption(
            "Technical workflow information "
            "for debugging and tracing."
        )

        col1, col2 = st.columns(2)

        with col1:
            st.write(
                "**Thread ID**"
            )
            st.code(
                st.session_state.thread_id
            )

            st.write(
                "**Research gaps**"
            )

            gaps = result.get(
                "research_gaps",
                [],
            )

            if gaps:
                for gap in gaps:
                    st.write(
                        f"- {gap}"
                    )
            else:
                st.caption(
                    "None"
                )

        with col2:
            st.write(
                "**Runtime errors**"
            )

            errors = result.get(
                "errors",
                [],
            )

            if errors:
                for error in errors:
                    st.error(error)
            else:
                st.caption(
                    "None"
                )

        st.write(
            "**Final workflow state**"
        )

        safe_state = {
            key: value
            for key, value in result.items()
            if key != "__interrupt__"
        }

        st.json(safe_state)


# -------------------------------------------------------------------
# Sidebar
# -------------------------------------------------------------------

with st.sidebar:

    st.title("Research Agent")

    st.caption(
        "Autonomous market research powered by "
        "LangGraph, tools, memory, guardrails, "
        "and human approval."
    )

    st.divider()

    st.markdown("### Workflow")

    workflow_items = [
        "Plan research",
        "Collect evidence",
        "Evaluate evidence",
        "Analyze findings",
        "Generate report",
        "Review and revise",
        "Human approval",
    ]

    for item in workflow_items:
        st.write(
            f"• {item}"
        )

    st.divider()

    st.markdown("### Agent capabilities")

    capabilities = [
        "Web research",
        "Market calculations",
        "Research memory",
        "Evidence evaluation",
        "Report revision",
        "Human approval",
    ]

    for capability in capabilities:
        st.write(
            f"• {capability}"
        )

    st.divider()

    st.caption(
        f"Session: "
        f"`{st.session_state.thread_id[:8]}…`"
    )

    if st.button(
        "Start a new research session",
        use_container_width=True,
    ):
        reset_session()
        st.rerun()


# -------------------------------------------------------------------
# Main header
# -------------------------------------------------------------------

st.title(
    "Autonomous Market Research Agent"
)

st.caption(
    "Turn a market question into a structured, "
    "evidence-based research report through an "
    "autonomous multi-step workflow."
)


# -------------------------------------------------------------------
# Idle / new research
# -------------------------------------------------------------------

if st.session_state.phase == "idle":

    st.divider()

    st.subheader(
        "What market would you like to research?"
    )

    st.caption(
        "Enter a market, industry, product category, "
        "or business domain. The agent will plan the "
        "research, collect evidence, analyze findings, "
        "and prepare a report for review."
    )

    with st.form(
        "research_form",
        clear_on_submit=False,
    ):

        topic = st.text_input(
            "Research topic",
            placeholder=(
                "e.g. Electronics market in Egypt"
            ),
            label_visibility="collapsed",
        )

        submitted = st.form_submit_button(
            "Start market research",
            type="primary",
            use_container_width=True,
        )

    if submitted:

        clean_topic = topic.strip()

        if not clean_topic:

            st.warning(
                "Enter a market research topic first."
            )

        else:

            with st.status(
                "Running market research...",
                expanded=True,
            ) as status:

                st.write(
                    "Planning the research scope"
                )

                st.write(
                    "Collecting and evaluating evidence"
                )

                st.write(
                    "Generating and reviewing the report"
                )

                run_research(
                    clean_topic
                )

                if (
                    st.session_state.phase
                    == "error"
                ):

                    status.update(
                        label="Research workflow failed",
                        state="error",
                        expanded=True,
                    )

                elif (
                    st.session_state.phase
                    == "approval"
                ):

                    status.update(
                        label=(
                            "Research completed — "
                            "review required"
                        ),
                        state="complete",
                        expanded=False,
                    )

                else:

                    status.update(
                        label="Research completed",
                        state="complete",
                        expanded=False,
                    )

            st.rerun()


# -------------------------------------------------------------------
# Human approval
# -------------------------------------------------------------------

elif st.session_state.phase == "approval":

    result = (
        st.session_state.result
        or {}
    )

    payload = (
        st.session_state.interrupt_payload
        or {}
    )

    st.info(
        payload.get(
            "message",
            "The research workflow is waiting "
            "for your approval.",
        )
    )

    st.markdown(
        f"### {st.session_state.topic}"
    )

    render_workflow(
        "approval",
        result,
    )

    render_status(
        result
    )

    report = choose_report(
        st.session_state.pending_report,
        payload.get("report"),
        result.get("draft_report"),
    )

    render_report(
        report,
        heading="Report ready for review",
    )

    st.divider()

    st.subheader(
        "Human review"
    )

    st.caption(
        "Review the generated report before the "
        "workflow is finalized."
    )

    approve_col, reject_col = st.columns(2)

    with approve_col:

        if st.button(
            "Approve report",
            type="primary",
            use_container_width=True,
        ):

            with st.spinner(
                "Finalizing approved report..."
            ):
                resume_after_human_decision(
                    "approve"
                )

            st.rerun()

    with reject_col:

        if st.button(
            "Reject report",
            use_container_width=True,
        ):

            with st.spinner(
                "Finalizing workflow..."
            ):
                resume_after_human_decision(
                    "reject"
                )

            st.rerun()


# -------------------------------------------------------------------
# Completed workflow
# -------------------------------------------------------------------

elif st.session_state.phase == "complete":

    result = (
        st.session_state.result
        or {}
    )

    approved = bool(
        result.get("approved")
    )

    if approved:

        st.success(
            "Research completed and the report "
            "was approved."
        )

    else:

        st.warning(
            "Research completed without human approval."
        )

    st.markdown(
        f"## {st.session_state.topic}"
    )

    render_workflow(
        "complete",
        result,
    )

    render_status(
        result
    )

    final_report = choose_report(
        st.session_state.final_report,
        st.session_state.pending_report,
        result.get("draft_report"),
    )

    render_report(
        final_report,
        heading="Final market research report",
    )

    render_memory(
        result
    )

    render_developer_details(
        result
    )

    st.divider()

    if st.button(
        "Research another market",
        type="primary",
        use_container_width=True,
    ):
        reset_session()
        st.rerun()


# -------------------------------------------------------------------
# Error
# -------------------------------------------------------------------

elif st.session_state.phase == "error":

    st.error(
        "The research workflow could not complete."
    )

    if st.session_state.error:

        with st.expander(
            "Technical error details",
            expanded=True,
        ):
            st.code(
                st.session_state.error
            )

    st.caption(
        "Check API keys, model availability, "
        "Tavily quota, and the terminal traceback."
    )

    if st.button(
        "Start a new session",
        type="primary",
        use_container_width=True,
    ):
        reset_session()
        st.rerun()